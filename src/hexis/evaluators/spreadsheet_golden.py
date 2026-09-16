# SPDX-License-Identifier: MIT AND CC-BY-SA-4.0
# The value normalization, cell comparison and range helpers and the LibreOffice recalculation command
# are adapted from SpreadsheetBench (https://github.com/RUCKBReasoning/SpreadsheetBench, CC BY-SA 4.0);
# see THIRD_PARTY_NOTICES.md.
"""Golden-workbook grading for SpreadsheetBench.

The check: compare the workbook saved by the agent with the golden workbook **value by value, cell by cell** within
the ``answer_position`` range declared by the task. One wrong cell means failure.

The comparison logic is taken verbatim from SpreadsheetBench's official ``evaluation.py`` (see
``third_party/SpreadsheetBench/evaluation.py``), including its type normalization: numbers rounded to two decimals,
dates converted to Excel serial numbers, ``""`` and ``None`` treated as equal. We do not write our own version: once
the criteria drift, the reported number is no longer SpreadsheetBench's number.

**The success rate cannot be gated on its absolute value.** Even the best models on SpreadsheetBench only reach about
half, so requiring a 90% absolute success rate from a compiled artifact is unattainable. The gate therefore asks a
**relative** question: with the same agent, the same model and the same verifier, does the set of tasks the compiled
artifact passes keep up with the original skill? See :func:`relative_success`.
"""

from __future__ import annotations

import datetime
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional

#: Location of the official dataset. Read-only, not committed to git (see .gitignore).
BENCH_ROOT = Path("third_party/SpreadsheetBench/spreadsheetbench_verified_400")


class GoldenUnavailable(RuntimeError):
    """The golden data is not available. **No degraded mode**: degraded grading would not follow SpreadsheetBench's criteria."""


# --------------------------------------------------------------------------- #
# Value comparison: taken verbatim from SpreadsheetBench's evaluation.py (CC BY-SA 4.0, see THIRD_PARTY_NOTICES.md)
# --------------------------------------------------------------------------- #
def _datetime_to_float(dt: datetime.datetime) -> float:
    excel_start_date = datetime.datetime(1899, 12, 30)
    delta = dt - excel_start_date
    return delta.days + delta.seconds / 86400.0


def _transform_value(v: Any) -> Any:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = round(float(v), 2)
    elif isinstance(v, datetime.time):
        v = str(v)[:-3]
    elif isinstance(v, datetime.datetime):
        v = round(_datetime_to_float(v), 0)
    elif isinstance(v, str):
        try:
            v = round(float(v), 2)
        except ValueError:
            pass
    return v


def compare_cell_value(v1: Any, v2: Any) -> bool:
    """Whether two cell values are the same. An empty string and ``None`` count as equal; that is the official rule."""
    v1, v2 = _transform_value(v1), _transform_value(v2)
    if (v1 == "" and v2 is None) or (v1 is None and v2 == ""):
        return True
    if (v1 == "" and v2 == "") or (v1 is None and v2 is None):
        return True
    if type(v1) is not type(v2):
        return False
    return v1 == v2


def _col_num2name(n: int) -> str:
    name = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        name = chr(65 + rem) + name
    return name


def _col_name2num(name: str) -> int:
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


def _parse_cell_range(range_str: str) -> tuple[tuple[int, int], tuple[int, int]]:
    start_cell, end_cell = range_str.split(":")
    sc = "".join(ch for ch in start_cell if not ch.isdigit())
    sr = "".join(ch for ch in start_cell if ch.isdigit())
    ec = "".join(ch for ch in end_cell if not ch.isdigit())
    er = "".join(ch for ch in end_cell if ch.isdigit())
    return (_col_name2num(sc), int(sr)), (_col_name2num(ec), int(er))


def cell_names(range_str: str) -> list[str]:
    """``'G2:G16'`` → ``['G2', ..., 'G16']``. A single cell is returned as itself."""
    if ":" not in range_str:
        return [range_str]
    (sc, sr), (ec, er) = _parse_cell_range(range_str)
    cols = [_col_num2name(i) for i in range(sc, ec + 1)]
    return [f"{c}{r}" for c in cols for r in range(sr, er + 1)]


# --------------------------------------------------------------------------- #
# Recalculation: openpyxl does not compute formulas, so without recalculation only None can be read
# --------------------------------------------------------------------------- #
#: LibreOffice executables. A prerequisite of SpreadsheetBench's official evaluation:
#: ``brew install --cask libreoffice`` / ``apt install libreoffice-calc``.
_SOFFICE_CANDIDATES = (
    "libreoffice", "soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice", "/usr/bin/libreoffice",
)


def find_soffice() -> Optional[str]:
    import shutil as _sh

    for c in _SOFFICE_CANDIDATES:
        hit = _sh.which(c)
        if hit:
            return hit
        if Path(c).is_file():
            return c
    return None


def recalculate(path: Path, *, timeout_s: float = 180.0) -> tuple[bool, str]:
    """Open the file headless in LibreOffice and save it back, so that formula results are stored as cached values.

    **This step is not optional.** openpyxl only writes formula strings and never computes them; reading back with
    ``data_only=True`` gives ``None``. Without this step a completely correct answer is graded as a failure: in a
    measured baseline-arm run, ``sb_49801`` wrote ``=LEFT(A1,2)&RIGHT(A1,5)`` (correct), yet grading reported
    "golden 'SL-0035', got None", and the success rate over eight tasks dropped from its real level to 2/8. The
    failures all had the shape "got None" rather than "computed wrong", which is exactly this bug's fingerprint:
    ``None`` is not a wrong answer, it is **no value at all**.

    The procedure and the invocation arguments are taken verbatim from the official ``evaluation/open_spreadsheet.py``
    (``--headless --calc --convert-to xlsx:Calc MS Excel 2007 XML``); we do not invent our own.

    Returns ``(success, reason)``. **Failure is never silent**: the caller must turn it into an explicit grading
    failure instead of falling back to comparing unrecalculated values, which would reintroduce this very bug.
    """
    import shutil as _sh
    import subprocess
    import tempfile as _tf

    soffice = find_soffice()
    if soffice is None:
        return False, ("LibreOffice not found, cannot recalculate formulas; "
                       "install it with brew install --cask libreoffice (macOS) / "
                       "apt install libreoffice-calc (Linux). "
                       "Without recalculation every formula answer is read as None and grading is systematically too low")
    path = Path(path).resolve()
    if not path.is_file():
        return False, f"file does not exist: {path}"
    with _tf.TemporaryDirectory() as tmp:
        try:
            proc = subprocess.run(
                [soffice, "--headless", "--calc",
                 "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                 "--outdir", tmp, str(path)],
                capture_output=True, text=True, timeout=timeout_s)
        except subprocess.TimeoutExpired:
            return False, f"recalculation timed out ({timeout_s:.0f}s)"
        out = Path(tmp) / (path.stem + ".xlsx")
        if proc.returncode != 0 or not out.is_file():
            return False, f"LibreOffice exit code {proc.returncode}: {(proc.stderr or '')[-200:]}"
        _sh.move(str(out), str(path.with_suffix(".xlsx")))
    return True, ""


def compare_workbooks(golden: Path, produced: Path, answer_position: str,
                      *, recalc: bool = True) -> tuple[bool, str]:
    """Compare values cell by cell within the ``answer_position`` range.

    ``answer_position`` looks like ``G2:G16`` or ``'Sheet2'!A1:B9``, with several comma-separated pieces; all of them
    must match to pass. ``data_only=True`` reads **cached values**, not formula strings: the task asks whether the
    result is right, and which formula computed it is outside the scope of grading.
    """
    import openpyxl

    if not Path(produced).is_file():
        return False, "no workbook produced"
    if recalc:
        ok, why = recalculate(Path(produced))
        if not ok:
            # recalculation impossible: fail and **state the reason**; never compare unrecalculated values instead.
            return False, f"cannot recalculate formulas: {why}"
        produced = Path(produced).with_suffix(".xlsx")
    try:
        wb_gt = openpyxl.load_workbook(filename=str(golden), data_only=True)
        wb_pr = openpyxl.load_workbook(filename=str(produced), data_only=True)
    except Exception as exc:  # noqa: BLE001 — a workbook that will not open fails; it must not crash the batch
        return False, f"{type(exc).__name__}: {exc}"

    for piece in answer_position.split(","):
        piece = piece.strip()
        if "!" in piece:
            sheet_name, cell_range = piece.split("!")
        else:
            sheet_name, cell_range = wb_gt.sheetnames[0], piece
        sheet_name = sheet_name.strip().strip("'")
        cell_range = cell_range.strip().strip("'")
        if sheet_name not in wb_pr.sheetnames:
            return False, f"no worksheet {sheet_name!r} in the output"
        ws_gt, ws_pr = wb_gt[sheet_name], wb_pr[sheet_name]
        for name in cell_names(cell_range):
            a, b = ws_gt[name].value, ws_pr[name].value
            if not compare_cell_value(a, b):
                return False, f"{sheet_name}!{name} golden {a!r}, got {b!r}"
    return True, ""


# --------------------------------------------------------------------------- #
# Dataset
# --------------------------------------------------------------------------- #
@dataclass
class GoldenTask:
    """One SpreadsheetBench task: instruction, initial workbook, golden workbook, grading range."""

    id: str
    instruction: str
    answer_position: str
    instruction_type: str
    init: Path
    golden: Path

    def verify(self, produced: Path) -> tuple[bool, str]:
        return compare_workbooks(self.golden, produced, self.answer_position)


def load_dataset(root: Path = BENCH_ROOT) -> dict[str, GoldenTask]:
    """Index the whole verified-400 set by task id.

    Each task directory holds three test cases (prefixes ``1_``/``2_``/``3_``), and the official evaluation requires
    all three to pass. Only the first is taken here: the closed loop checks whether the compiled artifact keeps up with
    the original skill on **the same task**, and the three cases are three data sets for the same instruction; running
    the other two only triples the tokens without adding discriminating power. This must be stated wherever the numbers
    are reported.
    """
    root = Path(root)
    meta = root / "dataset.json"
    if not meta.is_file():
        raise GoldenUnavailable(
            f"{meta} not found; SpreadsheetBench data is not distributed with the repository: "
            "get data/spreadsheetbench_verified_400.tar.gz from https://github.com/RUCKBReasoning/SpreadsheetBench "
            "and extract it into third_party/SpreadsheetBench/")
    out: dict[str, GoldenTask] = {}
    for row in json.loads(meta.read_text("utf-8")):
        tid = str(row["id"])
        d = root / "spreadsheet" / tid
        init, golden = d / f"1_{tid}_init.xlsx", d / f"1_{tid}_golden.xlsx"
        if not (init.is_file() and golden.is_file()):
            continue
        out[tid] = GoldenTask(
            id=tid, instruction=row["instruction"],
            answer_position=row["answer_position"],
            instruction_type=row.get("instruction_type", ""),
            init=init, golden=golden)
    if not out:
        raise GoldenUnavailable(f"{root} contains no complete task")
    return out


# --------------------------------------------------------------------------- #
# Relative success rate
# --------------------------------------------------------------------------- #
@dataclass
class SuccessReport:
    """The result of one grading run, with a record kept for every task."""

    per_task: dict[str, dict] = field(default_factory=dict)
    oracle_passed: list[str] = field(default_factory=list)
    candidate_passed: list[str] = field(default_factory=list)

    @property
    def n(self) -> int:
        return len(self.per_task)

    @property
    def oracle_rate(self) -> Optional[float]:
        return len(self.oracle_passed) / self.n if self.n else None

    @property
    def candidate_rate(self) -> Optional[float]:
        return len(self.candidate_passed) / self.n if self.n else None

    def to_dict(self) -> dict:
        return {
            "metric": "SpreadsheetBench golden-workbook success",
            "n": self.n,
            "oracle_passed": sorted(self.oracle_passed),
            "candidate_passed": sorted(self.candidate_passed),
            "oracle_success_rate": self.oracle_rate,
            "candidate_success_rate": self.candidate_rate,
            "relative_success": relative_success(self),
            "regressions": sorted(set(self.oracle_passed) - set(self.candidate_passed)),
            "tasks": self.per_task,
        }


def relative_success(report: SuccessReport) -> Optional[float]:
    """The share of **the tasks the original skill can do** that the compiled artifact passes.

    The denominator is the number of tasks the original skill passed, not the total number of tasks. When the compiled
    artifact fails a task the original skill could not do either, that must not be charged to compilation: it is the
    difficulty of the benchmark itself, and even the best models on SpreadsheetBench only reach about half. The gate
    asks "after moving into code, does what used to work still work?".

    Returns ``None`` (not measurable) rather than 1.0 when the original skill passed no task at all. "No regression"
    and "nothing was measured" must be distinguishable in the numbers; an earlier version misread failed runs as
    perfect protection in exactly this kind of spot.
    """
    if not report.oracle_passed:
        return None
    kept = set(report.oracle_passed) & set(report.candidate_passed)
    return len(kept) / len(report.oracle_passed)


def find_produced(root: Path, *, expected: str = "", shipped: Iterable[Path] = ()) -> Optional[Path]:
    """Pick the workbook the agent saved out of the exported artifacts.

    **Pick by name, not by time.** The first version picked "the most recently modified .xlsx" and got it wrong
    right away: the skill's ``assets/`` held the initial workbooks of all eight tasks, and the sandbox copies the whole
    skill directory at start, so all eight copies had fresh mtimes and "most recently modified" among them was
    essentially random. In a measured run, sb_53449 picked another task's untouched initial file and grading reported
    ``G4 golden None, got 'd'``: the agent had done nothing wrong; grading had compared the wrong file.

    So the order is: the output name the task specifies → the newest file left after excluding files **byte-identical**
    to the initial workbooks shipped with the skill → otherwise ``None``, which the caller grades as not passed.
    """
    root = Path(root)
    if not root.is_dir():
        return None
    xlsx = [f for f in root.rglob("*.xlsx") if not f.name.startswith("~$")]
    if not xlsx:
        return None
    if expected:
        hit = [f for f in xlsx if f.name == expected]
        if hit:
            return hit[0]
    # drop the initial workbooks shipped with the skill, byte for byte. One modified in place has new bytes and stays.
    seen = {f.read_bytes() for f in shipped if Path(f).is_file()}
    fresh = [f for f in xlsx if f.read_bytes() not in seen]
    if not fresh:
        return None
    return max(fresh, key=lambda f: f.stat().st_mtime)
